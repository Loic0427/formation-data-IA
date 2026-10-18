"""Génère un quiz Kahoot SQL progressif (format d'import officiel) :
- Kahoot_SQL_Zero_a_Expert_60Questions.xlsx : 60 questions en 4 paliers de 15
    1. Débutant      : SELECT, WHERE, ORDER BY, INSERT/UPDATE/DELETE, LIKE, NULL, COUNT
    2. Intermédiaire : GROUP BY/HAVING, jointures, clés, CASE, COALESCE, sous-requêtes
    3. Avancé        : CTE, fonctions de fenêtre, EXISTS, corrélées, récursif, ordre d'évaluation
    4. Expert        : ACID, index, sargabilité, EXPLAIN, isolation, MVCC, upsert, partitions

Usage : uv run --with openpyxl python build_kahoot_sql.py
Chaque question est écrite avec la BONNE réponse en premier ; le script mélange l'ordre (graine fixe)
et contrôle les limites Kahoot : question ≤ 120 caractères, réponse ≤ 75, durée autorisée.
Les questions Vrai/Faux (2 réponses) ne sont pas mélangées : la bonne réponse y est variée à la main.
"""
from pathlib import Path
import random
import shutil

import openpyxl

HERE = Path(__file__).parent
TEMPLATE = HERE / "KahootQuizTemplate.xlsx"
DUREES = {5, 10, 20, 30, 60, 90, 120, 240}

DEBUTANT = [
    ("SQL signifie…", ["Structured Query Language", "Simple Question Language", "Super Quick List", "Système de Questions Logiques"], 20),
    ("Quelle instruction lit des données dans une table ?", ["SELECT", "READ", "GET", "OPEN"], 20),
    ("Quelle clause indique la table à interroger ?", ["FROM", "WHERE", "TABLE", "IN"], 20),
    ("Que renvoie « SELECT * FROM clients » ?", ["Toutes les colonnes et lignes de clients", "La première ligne de clients", "Le nombre de clients", "Rien, * est interdit"], 20),
    ("Quelle clause filtre les lignes ?", ["WHERE", "FILTER", "HAVING", "ORDER BY"], 20),
    ("Quelle clause trie les résultats ?", ["ORDER BY", "SORT BY", "GROUP BY", "RANK BY"], 20),
    ("« ORDER BY prix DESC » trie…", ["Du plus cher au moins cher", "Du moins cher au plus cher", "Par ordre alphabétique", "Aléatoirement"], 20),
    ("Quelle instruction ajoute une ligne dans une table ?", ["INSERT INTO", "ADD ROW", "UPDATE", "CREATE"], 20),
    ("Quelle instruction modifie des lignes existantes ?", ["UPDATE", "MODIFY", "ALTER", "CHANGE"], 20),
    ("Que fait « DELETE FROM clients » sans WHERE ?", ["Supprime toutes les lignes", "Supprime la table", "Supprime une seule ligne", "Rien, erreur de syntaxe"], 30),
    ("Quel mot-clé limite le nombre de lignes renvoyées (PostgreSQL, MySQL) ?", ["LIMIT", "TOP", "FIRST", "MAX"], 20),
    ("Quel mot-clé supprime les doublons dans un SELECT ?", ["DISTINCT", "UNIQUE", "ONLY", "SINGLE"], 20),
    ("« WHERE email = NULL » permet de trouver les emails absents.", ["Faux", "Vrai"], 10),
    ("Que renvoie « WHERE nom LIKE 'Du%' » ?", ["Les noms commençant par Du", "Les noms contenant Du", "Les noms finissant par Du", "Les noms égaux à Du"], 30),
    ("Quelle fonction compte les lignes ?", ["COUNT", "SUM", "LEN", "NB"], 20),
]

INTERMEDIAIRE = [
    ("Quelle clause regroupe les lignes avant un agrégat (SUM, AVG…) ?", ["GROUP BY", "ORDER BY", "PARTITION", "CLUSTER"], 20),
    ("Pour filtrer sur un agrégat (ex. COUNT(*) > 5), on utilise…", ["HAVING", "WHERE", "FILTER", "LIMIT"], 30),
    ("Différence entre WHERE et HAVING ?", ["WHERE filtre avant agrégation, HAVING après", "Aucune, ce sont des synonymes", "HAVING filtre avant, WHERE après", "HAVING ne sert qu'aux tris"], 30),
    ("AVG(prix) ignore les valeurs NULL.", ["Vrai", "Faux"], 10),
    ("COUNT(*) et COUNT(colonne) renvoient…", ["Pas toujours la même chose (NULL ignorés)", "Toujours la même chose", "COUNT(*) est juste plus lent", "COUNT(colonne) compte les colonnes"], 30),
    ("Quelle jointure ne garde que les lignes présentes dans les deux tables ?", ["INNER JOIN", "LEFT JOIN", "FULL JOIN", "CROSS JOIN"], 20),
    ("« clients LEFT JOIN commandes » garde…", ["Tous les clients, même sans commande", "Seulement les clients avec commande", "Toutes les commandes sans client", "Seulement les commandes"], 30),
    ("Un CROSS JOIN entre 10 lignes et 20 lignes renvoie…", ["200 lignes", "30 lignes", "20 lignes", "10 lignes"], 30),
    ("Quelle clé identifie de façon unique chaque ligne d'une table ?", ["La clé primaire", "La clé étrangère", "La clé USB", "L'index"], 20),
    ("Une clé étrangère…", ["Référence la clé primaire d'une autre table", "Est toujours unique", "Est chiffrée", "Est un mot de passe"], 20),
    ("En SQL, que renvoie l'expression « NULL = NULL » ?", ["NULL (inconnu)", "TRUE", "FALSE", "Une erreur"], 30),
    ("CASE WHEN … THEN … END sert à…", ["Créer une valeur conditionnelle", "Changer la casse du texte", "Tester les majuscules", "Boucler sur les lignes"], 20),
    ("Que renvoie COALESCE(NULL, 'inconnu') ?", ["'inconnu'", "NULL", "Une erreur", "0"], 20),
    ("« BETWEEN '2024-01-01' AND '2024-01-31' » : les bornes sont…", ["Incluses toutes les deux", "Exclues toutes les deux", "Seule la première est incluse", "Seule la seconde est incluse"], 30),
    ("Que fait la sous-requête dans « WHERE id IN (SELECT id FROM vip) » ?", ["Elle fournit la liste des valeurs à filtrer", "Elle crée une table vip", "Elle trie les résultats", "Elle supprime les vip"], 30),
]

AVANCE = [
    ("Une CTE (table d'expression commune) commence par…", ["WITH", "CTE", "DECLARE", "TEMP"], 20),
    ("ROW_NUMBER() OVER (PARTITION BY client ORDER BY date) donne…", ["Un rang par client, dans l'ordre des dates", "Le nombre total de lignes", "Un rang global sur toute la table", "La date maximale par client"], 60),
    ("RANK() et DENSE_RANK() avec des ex æquo : la différence ?", ["RANK saute des rangs, DENSE_RANK non", "Ils sont identiques", "DENSE_RANK saute des rangs", "RANK ignore les ex æquo"], 60),
    ("LAG(montant) OVER (ORDER BY mois) renvoie…", ["Le montant du mois précédent", "Le montant du mois suivant", "La somme des montants", "Le montant maximal"], 30),
    ("SUM(montant) OVER (ORDER BY date) calcule…", ["Un cumul (somme courante)", "La somme totale sur chaque ligne", "La somme par date", "Une erreur de syntaxe"], 30),
    ("Fonction de fenêtre vs GROUP BY : la différence clé ?", ["La fenêtre garde toutes les lignes", "GROUP BY garde toutes les lignes", "Aucune différence", "La fenêtre est toujours plus lente"], 60),
    ("« WHERE EXISTS (SELECT 1 FROM commandes c WHERE c.client_id = cl.id) » teste…", ["Si au moins une commande existe pour le client", "Si la valeur 1 existe", "Si la table commandes existe", "Si client_id est unique"], 30),
    ("Une sous-requête corrélée…", ["Dépend de la ligne de la requête externe", "S'exécute une seule fois", "Est toujours plus rapide", "Est interdite dans WHERE"], 60),
    ("Une CTE récursive (WITH RECURSIVE) sert à…", ["Parcourir une hiérarchie (arbre, graphe)", "Accélérer les jointures", "Chiffrer les données", "Créer un index"], 30),
    ("Ordre d'évaluation logique d'une requête ?", ["FROM, WHERE, GROUP BY, HAVING, SELECT, ORDER BY", "SELECT, FROM, WHERE, GROUP BY, HAVING", "WHERE, FROM, SELECT, GROUP BY", "ORDER BY, SELECT, FROM, WHERE"], 60),
    ("UNION vs UNION ALL ?", ["UNION ALL garde les doublons et est plus rapide", "UNION ALL supprime les doublons", "Ils sont identiques", "UNION est plus rapide"], 30),
    ("NTILE(4) OVER (ORDER BY ca) fait…", ["Répartit les lignes en 4 quartiles", "Garde les 4 premières lignes", "Multiplie le ca par 4", "Calcule 4 moyennes"], 30),
    ("STRING_AGG (PostgreSQL) ou GROUP_CONCAT (MySQL)…", ["Concatène les valeurs d'un groupe en une chaîne", "Coupe une chaîne en morceaux", "Compte les chaînes", "Trie les chaînes"], 30),
    ("Un self join sert à…", ["Joindre une table avec elle-même (ex. employé → manager)", "Joindre sans clé", "Joindre deux bases différentes", "Éviter les doublons"], 30),
    ("FULL OUTER JOIN renvoie…", ["Toutes les lignes des deux tables, NULL si pas de correspondance", "Seules les correspondances", "Seulement la table de gauche", "Le produit cartésien"], 30),
]

EXPERT = [
    ("Dans ACID, le A signifie…", ["Atomicité", "Agrégation", "Asynchrone", "Autorisation"], 20),
    ("Un index B-tree accélère surtout…", ["Les recherches d'égalité et de plage", "Les INSERT massifs", "Les SELECT *", "Les LIKE '%mot%'"], 60),
    ("« WHERE LOWER(email) = 'a@b.fr' » n'utilise pas l'index sur email car…", ["La fonction rend le prédicat non sargable", "LOWER est interdit dans WHERE", "Les index ne gèrent pas le texte", "Il manque un LIMIT"], 60),
    ("Quelle commande montre le plan d'exécution réel (PostgreSQL) ?", ["EXPLAIN ANALYZE", "SHOW PLAN", "DESCRIBE QUERY", "TRACE"], 30),
    ("Un « Seq Scan » sur 10 M de lignes pour « WHERE id = 5 » indique…", ["Un index manquant ou inutilisé", "Une requête optimale", "Un problème réseau", "Un verrou sur la table"], 60),
    ("Selon la norme SQL, quel niveau d'isolation élimine les lectures fantômes ?", ["SERIALIZABLE", "READ COMMITTED", "READ UNCOMMITTED", "REPEATABLE READ"], 60),
    ("Un deadlock (interblocage), c'est…", ["Deux transactions qui s'attendent mutuellement", "Une table verrouillée pour toujours", "Un index corrompu", "Une requête sans WHERE"], 30),
    ("La bonne parade contre l'injection SQL ?", ["Les requêtes paramétrées (prepared statements)", "Mettre les mots-clés en majuscules", "Interdire les apostrophes", "Utiliser un VPN"], 30),
    ("Une vue matérialisée…", ["Stocke physiquement le résultat, à rafraîchir", "Est recalculée à chaque lecture", "Est une table temporaire", "Ne peut pas être indexée"], 60),
    ("« WHERE id NOT IN (SELECT id FROM t) » ne renvoie rien si la sous-requête contient un NULL.", ["Vrai", "Faux"], 30),
    ("ROLLUP, CUBE et GROUPING SETS servent à…", ["Produire des sous-totaux à plusieurs niveaux en une requête", "Créer des cubes OLAP physiques", "Arrondir les agrégats", "Trier les groupes"], 30),
    ("Le MVCC (PostgreSQL) permet…", ["Aux lecteurs et écrivains de ne pas se bloquer", "De compresser les tables", "De répliquer les données", "De chiffrer les colonnes"], 60),
    ("« INSERT … ON CONFLICT (id) DO UPDATE », c'est…", ["Un upsert", "Un rollback", "Une jointure", "Un trigger"], 30),
    ("Un index couvrant (covering index)…", ["Contient toutes les colonnes utiles : la table n'est pas lue", "Indexe toute la table", "Est un index chiffré", "Est forcément unique"], 60),
    ("Partitionner une grosse table par date sert surtout à…", ["Ignorer les partitions inutiles (partition pruning)", "Chiffrer par mois", "Trier par date", "Doubler l'espace disque"], 30),
]

SQL = DEBUTANT + INTERMEDIAIRE + AVANCE + EXPERT


def construire(questions, nom, graine, nb):
    assert len(questions) == nb, (nom, len(questions))
    rnd = random.Random(graine)
    dest = HERE / nom
    shutil.copy(TEMPLATE, dest)
    wb = openpyxl.load_workbook(dest)
    ws = wb.active
    positions = []
    for i, (q, reps, duree) in enumerate(questions):
        assert len(q) <= 120, (q, len(q))
        assert all(len(r) <= 75 for r in reps), [(r, len(r)) for r in reps if len(r) > 75]
        assert duree in DUREES
        bonne = reps[0]
        ordre = reps[:]
        if len(ordre) == 4:
            rnd.shuffle(ordre)
        r = 9 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, q)
        for j in range(4):
            ws.cell(r, 3 + j, ordre[j] if j < len(ordre) else None)
        ws.cell(r, 7, duree)
        ws.cell(r, 8, str(ordre.index(bonne) + 1))
        positions.append(ordre.index(bonne) + 1)
    wb.save(dest)
    print(f"{nom} : {nb} questions · bonnes réponses en position 1/2/3/4 : "
          f"{[positions.count(k) for k in (1, 2, 3, 4)]}")


if __name__ == "__main__":
    for bloc, nom in ((DEBUTANT, "débutant"), (INTERMEDIAIRE, "intermédiaire"), (AVANCE, "avancé"), (EXPERT, "expert")):
        assert len(bloc) == 15, (nom, len(bloc))
    construire(SQL, "Kahoot_SQL_Zero_a_Expert_60Questions.xlsx", 55, nb=60)
